"""Editable PDF -> PPTX conversion.

Instead of rendering each page to an image (the old behaviour), every page is
reconstructed from extracted content: one editable textbox per PDF text block
(runs keep font, size, bold/italic, color), real PowerPoint bullet/number
formatting for detected list lines, and images placed at their source
positions. Pages with no extractable text (scans) fall back to a full-page
image so nothing disappears.
"""

import io

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.features.pdf_prescan import parse_marker, prescan_pdf, normalize

EMU_PER_PT = 12700

_BOLD_FLAG = 1 << 4
_ITALIC_FLAG = 1 << 1

# PPT buAutoNum types by (format, suffix-style)
_AUTONUM = {
    ('decimal', '.'): 'arabicPeriod',
    ('decimal', ')'): 'arabicParenR',
    ('lowerLetter', '.'): 'alphaLcPeriod',
    ('lowerLetter', ')'): 'alphaLcParenR',
    ('upperLetter', '.'): 'alphaUcPeriod',
    ('upperLetter', ')'): 'alphaUcParenR',
    ('lowerRoman', '.'): 'romanLcPeriod',
    ('lowerRoman', ')'): 'romanLcParenR',
    ('upperRoman', '.'): 'romanUcPeriod',
    ('upperRoman', ')'): 'romanUcParenR',
}

_INDENT_STEP = 228600  # 0.25" in EMU


def _pt_emu(v):
    return int(v * EMU_PER_PT)


def _clean_font_name(name):
    if not name:
        return None
    # Strip PDF subset prefix like "ABCDEE+Calibri"
    if len(name) > 7 and name[6] == '+' and name[:6].isupper():
        name = name[7:]
    # Strip style suffixes pptx handles via bold/italic flags
    for suffix in ('-Bold', '-Italic', '-BoldItalic', ',Bold', ',Italic'):
        if name.endswith(suffix):
            name = name[: -len(suffix)]
    return name or None


def _span_runs(par, spans, skip_chars=0):
    """Emit runs for spans, skipping the first skip_chars characters."""
    remaining = skip_chars
    for span in spans:
        text = span.get('text', '')
        if remaining > 0:
            if len(text) <= remaining:
                remaining -= len(text)
                continue
            text = text[remaining:]
            remaining = 0
        if not text:
            continue
        run = par.add_run()
        run.text = text
        font = run.font
        size = span.get('size')
        if size:
            font.size = Pt(max(1, round(size * 2) / 2))
        flags = span.get('flags', 0)
        font.bold = bool(flags & _BOLD_FLAG)
        font.italic = bool(flags & _ITALIC_FLAG)
        name = _clean_font_name(span.get('font'))
        if name:
            font.name = name
        color = span.get('color')
        if color is not None:
            font.color.rgb = RGBColor(
                (color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)


def _get_or_add_ppr(paragraph):
    p = paragraph._p
    ppr = p.find(qn('a:pPr'))
    if ppr is None:
        ppr = p.makeelement(qn('a:pPr'), {})
        p.insert(0, ppr)
    return ppr


def _clear_bullet(ppr):
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buFont'):
        el = ppr.find(qn(tag))
        if el is not None:
            ppr.remove(el)


def _bullet_insert_index(ppr):
    # Schema: bullet props come after lnSpc/spcBef/spcAft/buClr/buSzTx...,
    # appending at the end of pPr is accepted by PowerPoint for bu* elements
    # as long as relative order buFont -> buChar holds.
    return len(ppr)


def _set_no_bullet(paragraph):
    ppr = _get_or_add_ppr(paragraph)
    _clear_bullet(ppr)
    ppr.append(ppr.makeelement(qn('a:buNone'), {}))


def _set_bullet_char(paragraph, glyph, level=0):
    ppr = _get_or_add_ppr(paragraph)
    _clear_bullet(ppr)
    if level:
        ppr.set('lvl', str(min(8, level)))
    ppr.set('marL', str(_INDENT_STEP * (level + 1)))
    ppr.set('indent', str(-_INDENT_STEP))
    bufont = ppr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    buchar = ppr.makeelement(qn('a:buChar'), {'char': glyph})
    ppr.append(bufont)
    ppr.append(buchar)


def _set_auto_num(paragraph, autonum_type, start_at, level=0):
    ppr = _get_or_add_ppr(paragraph)
    _clear_bullet(ppr)
    if level:
        ppr.set('lvl', str(min(8, level)))
    ppr.set('marL', str(_INDENT_STEP * (level + 1)))
    ppr.set('indent', str(-_INDENT_STEP))
    attrs = {'type': autonum_type}
    if start_at and start_at > 1:
        attrs['startAt'] = str(start_at)
    ppr.append(ppr.makeelement(qn('a:buAutoNum'), attrs))


_BULLET_GLYPH_MAP = {'◦': '○', '‣': '•', '·': '•', '▸': '▪', '►': '➢'}
_SAFE_BULLETS = set('•○▪■◆➢–—✓★')


def _bullet_glyph(marker):
    if not marker:
        return '•'
    ch = _BULLET_GLYPH_MAP.get(marker, marker)
    return ch if ch in _SAFE_BULLETS else '•'


class _LogicalLine:
    __slots__ = ('spans', 'x0', 'y0', 'y1', 'size')

    def __init__(self, line):
        self.spans = list(line.get('spans', []))
        bbox = line.get('bbox', (0, 0, 0, 0))
        self.x0, self.y0 = bbox[0], bbox[1]
        self.y1 = bbox[3]
        sizes = [s.get('size', 11) for s in self.spans if s.get('text', '').strip()]
        self.size = max(sizes) if sizes else 11

    @property
    def text(self):
        return ''.join(s.get('text', '') for s in self.spans)


def _merge_wrapped_lines(lines):
    """Group physical lines into logical paragraphs (soft-wrap reflow)."""
    if not lines:
        return []
    groups = [[lines[0]]]
    for prev, cur in zip(lines, lines[1:]):
        gap = cur.y0 - prev.y1
        marker, _, _ = parse_marker(cur.text)
        new_par = (
            marker is not None
            or gap > 0.45 * cur.size
            or cur.x0 < prev.x0 - 2  # outdent = new paragraph
        )
        if new_par:
            groups.append([cur])
        else:
            groups[-1].append(cur)
    return groups


def _marker_ordinal(marker, kind):
    """(family, ordinal, style) for simple single-part sequence markers."""
    body = marker.rstrip('.)')
    style = ')' if marker.endswith(')') else '.'
    if kind == 'decimal' and body.isdigit():
        return ('decimal', int(body), style)
    if kind == 'alpha' and len(body) == 1:
        fam = 'lowerLetter' if body.islower() else 'upperLetter'
        return (fam, ord(body.lower()) - 96, style)
    if kind == 'roman':
        from app.features.docx_list_fixer import _roman_to_int
        val = _roman_to_int(body)
        if val:
            fam = 'lowerRoman' if body.islower() else 'upperRoman'
            return (fam, val, style)
    return None


def _add_text_block(slide, block, scan):
    bbox = block.get('bbox', (0, 0, 100, 20))
    x0, y0, x1, y1 = bbox
    box = slide.shapes.add_textbox(
        _pt_emu(x0), _pt_emu(y0),
        max(_pt_emu(x1 - x0 + 3), _pt_emu(8)),
        max(_pt_emu(y1 - y0 + 2), _pt_emu(8)),
    )
    tf = box.text_frame
    tf.word_wrap = True
    for attr in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, attr, 0)

    lines = [_LogicalLine(l) for l in block.get('lines', [])
             if any(s.get('text', '').strip() for s in l.get('spans', []))]
    groups = _merge_wrapped_lines(lines)
    if not groups:
        return

    # First pass: parse markers per paragraph group.
    parsed = []
    for group in groups:
        text = ''.join(l.text for l in group)
        marker, kind, _ = parse_marker(text)
        info = None
        if marker is None and scan is not None:
            info = scan.by_key.get(normalize(text.strip())[:40])
            if info is not None and info.kind == 'vector-bullet':
                marker, kind = '•', 'vector-bullet'
        parsed.append((group, text, marker, kind))

    # Decide auto-numbering per family: only when the block's sequence is
    # contiguous from its first item (exactness guarantee, PPT flavour).
    seq_state = {}
    seq_ok = {}
    for _, _, marker, kind in parsed:
        if marker is None or kind in ('bullet', 'vector-bullet', 'hybrid',
                                      'multilevel', None):
            continue
        ord_info = _marker_ordinal(marker, kind)
        if ord_info is None:
            continue
        fam, val, style = ord_info
        key = (fam, style)
        if key not in seq_state:
            seq_state[key] = val
            seq_ok[key] = True
        else:
            seq_state[key] += 1
            if seq_state[key] != val:
                seq_ok[key] = False

    seq_started = set()
    first_par = True
    for group, text, marker, kind in parsed:
        par = tf.paragraphs[0] if first_par else tf.add_paragraph()
        first_par = False

        strip = 0
        spans = [s for l in group for s in l.spans]
        full = ''.join(s.get('text', '') for s in spans)
        lstrip_n = len(full) - len(full.lstrip())

        if marker is not None and kind in ('bullet',):
            strip = lstrip_n + len(marker)
            rest = full[strip:]
            strip += len(rest) - len(rest.lstrip(' \t'))
            _span_runs(par, spans, skip_chars=strip)
            _set_bullet_char(par, _bullet_glyph(marker))
        elif kind == 'vector-bullet':
            _span_runs(par, spans, skip_chars=lstrip_n)
            _set_bullet_char(par, '•')
        elif marker is not None and kind in ('decimal', 'alpha', 'roman'):
            ord_info = _marker_ordinal(marker, kind)
            key = (ord_info[0], ord_info[2]) if ord_info else None
            autonum = _AUTONUM.get((ord_info[0], ord_info[2])) if ord_info else None
            if ord_info and autonum and seq_ok.get(key):
                strip = lstrip_n + len(marker)
                rest = full[strip:]
                strip += len(rest) - len(rest.lstrip(' \t'))
                _span_runs(par, spans, skip_chars=strip)
                start = None
                if key not in seq_started:
                    seq_started.add(key)
                    # startAt = first literal ordinal in this textbox
                    first_val = None
                    for _, _, m2, k2 in parsed:
                        if m2 and k2 == kind:
                            oi = _marker_ordinal(m2, k2)
                            if oi and (oi[0], oi[2]) == key:
                                first_val = oi[1]
                                break
                    start = first_val
                _set_auto_num(par, autonum, start)
            else:
                # Non-contiguous or exotic: keep literal text (exact).
                _span_runs(par, spans, skip_chars=lstrip_n)
                _set_no_bullet(par)
        elif marker is not None and kind in ('multilevel', 'hybrid'):
            # Literal marker text with hanging indent look.
            level = max(0, marker.count('.') - (0 if marker.endswith('.') else -1) - 1)
            _span_runs(par, spans, skip_chars=lstrip_n)
            ppr = _get_or_add_ppr(par)
            _clear_bullet(ppr)
            ppr.set('marL', str(_INDENT_STEP * (min(level, 8) + 1)))
            ppr.set('indent', str(-_INDENT_STEP))
            ppr.append(ppr.makeelement(qn('a:buNone'), {}))
        else:
            _span_runs(par, spans, skip_chars=lstrip_n)
            _set_no_bullet(par)


def _add_images(slide, page, doc):
    try:
        infos = page.get_image_info(xrefs=True)
    except Exception:
        return
    seen = set()
    for info in infos:
        xref = info.get('xref', 0)
        bbox = info.get('bbox')
        if not xref or bbox is None:
            continue
        key = (xref, tuple(round(v, 1) for v in bbox))
        if key in seen:
            continue
        seen.add(key)
        x0, y0, x1, y1 = bbox
        if (x1 - x0) < 8 or (y1 - y0) < 8:
            continue
        try:
            img = doc.extract_image(xref)
            stream = io.BytesIO(img['image'])
            slide.shapes.add_picture(
                stream, _pt_emu(x0), _pt_emu(y0),
                _pt_emu(x1 - x0), _pt_emu(y1 - y0))
        except Exception:
            continue


def _add_fullpage_image(slide, page, slide_w, slide_h):
    mat = fitz.Matrix(300 / 72, 300 / 72)
    pix = page.get_pixmap(matrix=mat)
    stream = io.BytesIO(pix.tobytes('png'))
    slide.shapes.add_picture(stream, Emu(0), Emu(0), slide_w, slide_h)


def convert(input_path, out_path):
    doc = fitz.open(input_path)
    try:
        scan = prescan_pdf(input_path)
        prs = Presentation()
        for page in doc:
            slide_w = Emu(_pt_emu(page.rect.width))
            slide_h = Emu(_pt_emu(page.rect.height))
            prs.slide_width = slide_w
            prs.slide_height = slide_h
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            data = page.get_text('dict')
            text_blocks = [b for b in data.get('blocks', [])
                           if b.get('type') == 0 and any(
                               s.get('text', '').strip()
                               for l in b.get('lines', [])
                               for s in l.get('spans', []))]

            if not text_blocks:
                _add_fullpage_image(slide, page, slide_w, slide_h)
                continue

            _add_images(slide, page, doc)
            for block in sorted(text_blocks,
                                key=lambda b: (round(b['bbox'][1], 1), b['bbox'][0])):
                _add_text_block(slide, block, scan)

        prs.save(out_path)
    finally:
        doc.close()
    return out_path
