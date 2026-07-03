import io
import shutil
from pathlib import Path

import fitz
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse

from app.core.config import TEMP_DIR
from app.core.exceptions import ConversionError
from app.core.file_manager import cleanup_file, get_output_path, save_upload

router = APIRouter(prefix="/api/pdf-to-office", tags=["pdf-to-office"])

ALLOWED_FORMATS = {"docx", "pptx"}

MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def convert_pdf_to_docx(input_path: str) -> str:
    from pdf2docx import Converter

    out_path = str(get_output_path(".docx"))
    try:
        cv = Converter(input_path)
        cv.convert(out_path)
        cv.close()
    except Exception as e:
        raise ConversionError(f"PDF to DOCX conversion failed: {e}")

    if not Path(out_path).exists():
        raise ConversionError("Conversion produced no output file")
    return out_path


def convert_pdf_to_pptx(input_path: str) -> str:
    from pptx import Presentation
    from pptx.util import Emu

    try:
        doc = fitz.open(input_path)
    except Exception as e:
        raise ConversionError(f"Failed to open PDF: {e}")

    prs = Presentation()
    dpi = 300
    zoom = dpi / 72

    for page_num in range(len(doc)):
        page = doc[page_num]
        mat = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat)

        slide_width = Emu(int(page.rect.width * 914400 / 72))
        slide_height = Emu(int(page.rect.height * 914400 / 72))
        prs.slide_width = slide_width
        prs.slide_height = slide_height

        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)

        img_bytes = pix.tobytes("png")
        img_stream = io.BytesIO(img_bytes)
        slide.shapes.add_picture(img_stream, Emu(0), Emu(0), slide_width, slide_height)

    doc.close()

    out_path = str(get_output_path(".pptx"))
    prs.save(out_path)
    return out_path


@router.post("/")
async def pdf_to_office_endpoint(
    file: UploadFile = File(...),
    format: str = Form("docx"),
):
    fmt = format.lower().strip()
    if fmt not in ALLOWED_FORMATS:
        raise HTTPException(status_code=400, detail=f"Unsupported format: {fmt}. Use 'docx' or 'pptx'.")

    suffix = Path(file.filename).suffix.lower()
    if suffix != ".pdf":
        raise HTTPException(status_code=400, detail="Please upload a PDF file.")

    temp_path = save_upload(file)
    try:
        if fmt == "docx":
            out_path = convert_pdf_to_docx(str(temp_path))
        else:
            out_path = convert_pdf_to_pptx(str(temp_path))

        stem = Path(file.filename).stem
        return FileResponse(
            path=out_path,
            media_type=MEDIA_TYPES[fmt],
            filename=f"{stem}.{fmt}",
        )
    finally:
        cleanup_file(temp_path)
